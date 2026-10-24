# -*- coding: utf-8 -*-
"""把 Markdown 主稿渲染成带图的 Word 文档（中英文共用）。

【为什么需要这个脚本】
`03_主稿_中文投稿版.docx` 是 2026-10-19 21:18 用临时手段生成的，之后 md 又经过
第 55–57 轮三轮修改（新增五组结果 + 27 张 SI 图引用），docx 已严重滞后。
本脚本以 slug 化的 md 为唯一输入，重新生成可交付的 Word 稿。

【设计要点】
1. **单一输入**：只读 md，不读旧 docx。旧 docx 仅用于比对样式基线（A4、2.5 cm 边距）。
2. **公式转原生对象**：md 里的 `$...$` LaTeX 经 `_latex_omml.py` 转成 OMML
   （Office Math），采用**线性排版**——分式写作 `a/b`，上下标是真上下标。
   Word/WPS 里双击即可编辑，不是死文本。早期版本印浅黄底 LaTeX 源码，
   作者明确要求改为线性格式，故重写。
3. **图位置**：按 md 中 `**图 N.**` / `**Figure N.**` 图注的位置插图，
   图宽固定 16 cm（正文宽 = 21 − 2×2.5 = 16 cm），图注紧随其下。
4. **表格**：解析 md 的管道表格，转成带边框的 Word 表格。
5. **样式**：A4、2.5 cm 边距、正文 Times New Roman 12 pt（中文用宋体）、
   行距 1.5、标题分级。

用法：
    python _build_docx.py --md 03_主稿_中文投稿版.md --out 03_主稿_中文投稿版_带图.docx --lang zh
    python _build_docx.py --md 04_wr_manuscript_en.md --out 04_wr_manuscript_en_带图.docx --lang en
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _latex_omml import latex_to_omath_reported  # noqa: E402

HERE = Path(__file__).resolve().parent
FIGDIR = HERE / "figs"

# 主图：图号 -> 图片文件（不含扩展名）
MAIN_FIGS = {
    1: "Fig1_framework",
    2: "Fig2_bottleneck_migration",
    3: "Fig3_constraint_trajectories",
    4: "Fig4_interventions",
    5: "Fig5_city_gradient_policy",
    6: "Fig6_timescale_reservoir",
}

BODY_FONT = "Times New Roman"
CJK_FONT = "宋体"
HEAD_FONT = "Times New Roman"
CJK_HEAD = "黑体"


# ---------------------------------------------------------------------------
# 样式
# ---------------------------------------------------------------------------
def _set_run_fonts(run, name: str, cjk: str, size: float | None = None,
                   bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for attr, val in (("left_margin", 2.5), ("right_margin", 2.5),
                      ("top_margin", 2.54), ("bottom_margin", 2.54)):
        setattr(sec, attr, Cm(val))

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)

    for name, size, cjk in (("Heading 1", 16, CJK_HEAD), ("Heading 2", 14, CJK_HEAD),
                            ("Heading 3", 12.5, CJK_HEAD), ("Heading 4", 12, CJK_HEAD)):
        st = doc.styles[name]
        st.font.name = HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
    return doc


# ---------------------------------------------------------------------------
# 行内格式：**粗体**、`代码`、$数学$
# ---------------------------------------------------------------------------
INLINE_RE = re.compile(
    r"(\*\*.+?\*\*"          # 粗体
    r"|(?<!\*)\*[^*\n]+?\*(?!\*)"   # 斜体（期刊名等）
    r"|`[^`]+`"              # 行内代码
    r"|\$[^$]+\$)"           # 数学
)

# 本次构建里公式转换的统计（未知 LaTeX 命令、失败次数）
MATH_STATS = {"n": 0, "unknown": {}, "failed": []}


def _shade(run, hex_fill: str) -> None:
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    rpr.append(shd)


def add_math(par, latex: str, size: float = 12) -> None:
    """把一段 LaTeX 作为**原生 OMML 公式**插入段落（线性排版）。

    转换失败时降级为普通文本（保留可读的线性写法），不吞内容。
    """
    try:
        xml, unknown = latex_to_omath_reported(latex, size)
        par._p.append(parse_xml(xml))
        MATH_STATS["n"] += 1
        for u in unknown:
            MATH_STATS["unknown"][u] = MATH_STATS["unknown"].get(u, 0) + 1
    except Exception as e:                      # noqa: BLE001
        MATH_STATS["failed"].append((latex, repr(e)))
        from _latex_omml import latex_to_plain
        r = par.add_run(latex_to_plain(latex) or latex)
        _set_run_fonts(r, "Cambria Math", CJK_FONT, size - 0.5)


def add_inline(par, text: str, size: float = 12, base_bold: bool = False) -> None:
    """把一段含行内标记的文本写入段落。"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            _set_run_fonts(r, BODY_FONT, CJK_FONT, size, base_bold)
        tok = m.group(0)
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2])
            _set_run_fonts(r, BODY_FONT, CJK_FONT, size, True)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            # 斜体：期刊名（*Water Research*、*ES&T*）等
            r = par.add_run(tok[1:-1])
            _set_run_fonts(r, BODY_FONT, CJK_FONT, size, base_bold)
            r.font.italic = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            _set_run_fonts(r, "Consolas", CJK_FONT, size - 1, base_bold)
            _shade(r, "F2F2F2")
        else:  # 数学：转原生 OMML 线性公式
            add_math(par, tok[1:-1], size)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        _set_run_fonts(r, BODY_FONT, CJK_FONT, size, base_bold)


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------
TABLE_FONT = 9.5


def _visual_len(s: str) -> int:
    """估算一段 md 文本在表格里占的宽度（CJK 记 2，其它记 1）。

    公式按线性形式后的可见长度估，避免 `$u=(\\text{再生水利用量})/...$`
    这类超长公式把整列撑爆。
    """
    from _latex_omml import latex_to_plain

    def width(t: str) -> int:
        w = 0
        for ch in t:
            w += 2 if ord(ch) > 0x2E80 else 1
        return w

    total = 0
    pos = 0
    for m in INLINE_RE.finditer(s):
        total += width(s[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("$"):
            total += width(latex_to_plain(tok[1:-1]))
        elif tok.startswith("**"):
            total += width(tok[2:-2])
        else:
            total += width(tok[1:-1])
        pos = m.end()
    total += width(s[pos:])
    return total


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # 按内容宽度分配列宽，总宽 = 正文宽 16 cm。
    # 术语表这类"符号 + 含义 + 单位"三列结构里，含义列天然长得多，
    # 若纯按字数比例分，符号列会被压到 `CAWCC` 折行。故对首列给一个
    # 与最长不可断词相关的下限。
    weights = []
    for j in range(ncol):
        mx = max((_visual_len(r[j]) if j < len(r) else 0) for r in rows)
        weights.append(max(mx, 6))
    total_w = sum(weights)
    usable = 16.0
    widths = [max(1.9, usable * w / total_w) for w in weights]

    # 首列（符号）不窄于其最长连续拉丁词宽 + 余量，避免被拆词
    first_max = max((_visual_len(r[0]) if r else 0) for r in rows)
    widths[0] = max(widths[0], min(4.0, 0.30 * first_max + 0.9))

    scale = usable / sum(widths)
    widths = [w * scale for w in widths]

    # 列宽必须同时写到 gridCol 与每个 tcW：python-docx 的 cell.width 只写
    # tcW，LibreOffice/WPS 渲染时优先读 gridCol，只写 tcW 会被忽略、表格
    # 退回等宽，长公式就会溢出单元格。
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)

    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.width = Cm(widths[j])
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            txt = row[j] if j < len(row) else ""
            # 单元格内允许用 <br> 手动换行（md 表格里换行不能用空行）
            for k, seg in enumerate(txt.split("<br>")):
                if k:
                    par.add_run().add_break()
                add_inline(par, seg.strip(), size=TABLE_FONT, base_bold=(i == 0))
            if i == 0 or j > 0:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 表头行跨页重复
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    # 小表（≤12 行）禁止跨页断行。
    # 注意 Word 的语义：cantSplit 只保证"单行不被拆开"，并不阻止整张表
    # 跨页；真正让整表落在同一页，必须把除最后一行外所有行的段落都设
    # keep_with_next（Word 的分页算法据此把行绑在一起）。
    if len(rows) <= 12:
        for row in table.rows:
            tr = row._tr
            pr = tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            pr.append(cs)
        for row in table.rows[:-1]:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True

    return


# ---------------------------------------------------------------------------
# 主渲染
# ---------------------------------------------------------------------------
def _set_list_start(par, start: int) -> None:
    """（已弃用，保留占位以免旧调用出错；见 add_numbered_para。）"""
    return


def _join_continuation(lines: list[str], i: int, head: str) -> tuple[str, int]:
    """把列表项的缩进续行并入同一段。

    Markdown 里列表项可以这样写（本项目文档大量使用）：

        1. **闭式阈值**：$r^{*}=...$ = 0.5532。
           三组分取最紧：TDS 0.5532 / 氯化物 0.6842。

    第二行是首项的续行，不是新段落。早期的渲染器把每一行都当成独立段落，
    Word 里就变成"编号列表项 + 缩进的散段"，读起来像排版错乱。
    这里把缩进续行用空格拼到首项后面，返回拼好的文本与新的行号。

    停止条件（任一满足即返回）：
      - 行耗尽 / 空行
      - 不以空白开头的行（新段落或新列表项）
      - 以 --- 开头的分隔线
      - 以 # 开头的标题
    """
    parts = [head]
    while i < len(lines):
        nxt = lines[i]
        stripped = nxt.strip()
        if not stripped:
            break
        if not nxt[:1].isspace():          # 无缩进 = 新块
            break
        if stripped.startswith(("- ", "* ", "• ")) or re.match(r"^\d+\.\s", stripped):
            break                          # 嵌套列表，暂不合并
        if re.fullmatch(r"-{3,}", stripped):
            break
        if stripped.startswith("#"):
            break
        parts.append(stripped)
        i += 1
    return " ".join(parts), i


def add_numbered_para(doc, number: int, body: str) -> None:
    """输出一个编号段落，编号与 md 源完全一致。

    不用 Word 的自动编号（"List Number" 样式引用同一个 numId，跨列表会连续
    编号：上一组 1–6，下一组就从 7 开始）。本项目一个文档里有多组小列表，
    自动编号会产生明显错误的序号。这里把编号写成文本前缀 + 悬挂缩进，
    视觉上等同于编号列表，且序号严格可控。
    """
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.left_indent = Cm(0.85)
    pf.first_line_indent = Cm(-0.85)
    pf.space_after = Pt(4)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = par.add_run(f"{number}. ")
    _set_run_fonts(r, BODY_FONT, CJK_FONT, 12, False)
    add_inline(par, body)


def render(md_path: Path, out_path: Path, lang: str) -> dict:
    text = md_path.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = setup_document()
    stats = {"paras": 0, "tables": 0, "figs": 0, "fig_missing": []}

    fig_caption_re = (re.compile(r"^\*\*图\s*(\d+)\.\*\*\s*(.*)$") if lang == "zh"
                      else re.compile(r"^\*\*Figure\s*(\d+)\.\*\*\s*(.*)$"))

    i = 0
    pending_table: list[list[str]] = []
    in_block_math = False
    _pending_math: list[str] = []

    def flush_table() -> None:
        nonlocal pending_table
        if pending_table:
            # 表格与前一段（表题）绑定在同一页，避免表题落页尾、表体跑到次页
            if doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            add_table(doc, pending_table)
            stats["tables"] += 1
            pending_table = []

    while i < len(lines):
        line = lines[i].rstrip()

        # 围栏代码块 ```...```：等宽字体 + 浅灰底，保留原始换行
        if line.strip().startswith("```"):
            flush_table()
            i += 1
            block: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1                     # 跨过结束围栏
            for raw in block:
                cp = doc.add_paragraph()
                pf = cp.paragraph_format
                pf.space_after = Pt(0)
                pf.space_before = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                pf.left_indent = Cm(0.6)
                r = cp.add_run(raw if raw else " ")
                _set_run_fonts(r, "Consolas", CJK_FONT, 9.5)
                _shade(r, "F5F5F5")
                stats["paras"] += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 块级公式 $$...$$（可跨行）：独占一段，居中，不输出 $$ 记号
        if in_block_math or line.lstrip().startswith("$$"):
            flush_table()
            body = line.strip()
            started = body.startswith("$$")
            if started:
                body = body[2:]
            ends = body.endswith("$$")
            if ends:
                body = body[:-2]
            _pending_math.append(body)
            if ends:
                latex = "\n".join(_pending_math).strip()
                _pending_math.clear()
                in_block_math = False
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(6)
                par.paragraph_format.space_after = Pt(6)
                add_math(par, latex, size=12)
                stats["paras"] += 1
            else:
                in_block_math = True
            i += 1
            continue

        # 空行
        if not line.strip():
            flush_table()
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", line.strip()):
            flush_table()
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_table()
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(12)
                add_inline(par, title, size=18, base_bold=True)
            else:
                h = doc.add_heading(level=min(level - 1, 4))
                h.text = ""
                add_inline(h, title, size={2: 14, 3: 12.5}.get(level, 12), base_bold=True)
            i += 1
            continue

        # 管道表格
        if line.lstrip().startswith("|") and line.rstrip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            # 跳过分隔行 |---|---|
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                i += 1
                continue
            pending_table.append(cells)
            i += 1
            continue

        # 图注：紧接之后插入图片
        m = fig_caption_re.match(line)
        if m:
            flush_table()
            num = int(m.group(1))
            stem = MAIN_FIGS.get(num)
            img = FIGDIR / f"{stem}.png" if stem else None
            if img and img.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.add_run().add_picture(str(img), width=Cm(16.0))
                stats["figs"] += 1
            else:
                stats["fig_missing"].append(num)
            # 图注：**图 N.** 用粗体标签 + 正文，两段拼接。
            # 旧写法用切片推算前缀长度，实际把 "**图 1.**" 整段切掉，
            # 导致 Word 里图注从 "研究边界与…" 开始、丢了图号。
            label = f"图 {num}." if lang == "zh" else f"Figure {num}."
            body = m.group(2).strip()
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            cap.paragraph_format.space_after = Pt(10)
            r = cap.add_run(label + "  ")
            _set_run_fonts(r, BODY_FONT, CJK_FONT, 10, True)
            add_inline(cap, body, size=10)
            stats["paras"] += 1
            i += 1
            continue

        # 列表项
        m = re.match(r"^(\s*)[-*•]\s+(.*)$", line)
        if m:
            flush_table()
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(4)
            body, i = _join_continuation(lines, i + 1, m.group(2).strip())
            add_inline(par, body)
            stats["paras"] += 1
            continue

        # 编号列表
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            flush_table()
            body, i = _join_continuation(lines, i + 1, m.group(3).strip())
            add_numbered_para(doc, int(m.group(2)), body)
            stats["paras"] += 1
            continue

        # 引用块
        if line.lstrip().startswith(">"):
            flush_table()
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            add_inline(par, line.lstrip()[1:].strip(), size=10.5)
            stats["paras"] += 1
            i += 1
            continue

        # 普通段落
        flush_table()
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(par, line.strip())
        stats["paras"] += 1
        i += 1

    flush_table()
    doc.save(out_path)
    return stats


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--lang", default="zh", choices=("zh", "en"))
    args = ap.parse_args()

    md = Path(args.md)
    if not md.is_absolute():
        md = HERE / md
    out = Path(args.out)
    if not out.is_absolute():
        out = HERE / out
    st = render(md, out, args.lang)
    size_mb = out.stat().st_size / 1e6
    print(f"[OK] {out.name}  {size_mb:.2f} MB")
    print(f"     段落 {st['paras']}｜表格 {st['tables']}｜插图 {st['figs']}")
    print(f"     公式 {MATH_STATS['n']} 个 OMML 对象")
    if MATH_STATS["unknown"]:
        print(f"     ⚠ 未知 LaTeX 命令: {MATH_STATS['unknown']}")
    if MATH_STATS["failed"]:
        print(f"     ⚠ 公式转换失败 {len(MATH_STATS['failed'])} 处:")
        for tex, err in MATH_STATS["failed"][:5]:
            print(f"        {tex!r} -> {err}")
    if st["fig_missing"]:
        print(f"     ⚠ 缺图号: {st['fig_missing']}")


if __name__ == "__main__":
    main()

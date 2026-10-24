# -*- coding: utf-8 -*-
"""快速体检 PPT：页数、嵌入图片数、每页文字量。"""
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation
import zipfile

HERE = Path(__file__).resolve().parent
DECK = HERE / "62_导师汇报_10页_20261020.pptx"

prs = Presentation(str(DECK))
print("页数:", len(prs.slides))

z = zipfile.ZipFile(str(DECK))
imgs = [n for n in z.namelist() if n.startswith("ppt/media/")]
print("嵌入图片:", len(imgs))
for n in imgs:
    print("   ", n.split("/")[-1])

for i, s in enumerate(prs.slides, 1):
    n_txt = sum(1 for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    n_pic = sum(1 for sh in s.shapes if sh.shape_type == 13)
    n_shp = len(s.shapes)
    first = ""
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            first = sh.text_frame.text.strip().split("\n")[0][:34]
            break
    print(f"  P{i:02d}  形状{n_shp:3d}  文字框{n_txt:3d}  图{n_pic}  首行: {first}")

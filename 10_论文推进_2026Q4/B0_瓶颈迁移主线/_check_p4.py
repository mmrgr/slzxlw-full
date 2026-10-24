# -*- coding: utf-8 -*-
"""粗估每个文本框所需高度，找出可能溢出的（文字被截断/压到邻居）。

python-pptx 不做排版计算，文字溢出在生成时不报错，只有投出来才发现，
所以这里用"字符宽度 × 行数"粗估。判据放宽到 1.35 倍，只报明显溢出。
"""
from __future__ import annotations

import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation

HERE = Path(__file__).resolve().parent
DECK = HERE / "62_导师汇报_10页_20261020.pptx"

EMU_PER_PT = 12700


def est_line_count(text: str, width_emu: int, size_pt: float) -> int:
    """中文按 1.0 字宽、西文按 0.55 字宽累加，超过可用宽度即换行。"""
    usable_pt = width_emu / EMU_PER_PT
    if usable_pt <= 0:
        return 1
    lines, cur = 1, 0.0
    for ch in text:
        w = size_pt * (1.0 if ord(ch) > 0x2E80 else 0.55)
        cur += w
        if cur > usable_pt:
            lines += 1
            cur = w
    return lines


def main() -> int:
    prs = Presentation(str(DECK))
    bad = 0
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text.strip()
            if not txt:
                continue
            need_emu = 0
            maxsz = 12.0
            for p in sh.text_frame.paragraphs:
                ptxt = "".join(r.text for r in p.runs)
                if not ptxt:
                    continue
                size = next((r.font.size.pt for r in p.runs if r.font.size), 14.0)
                maxsz = max(maxsz, size)
                ls = p.line_spacing or 1.25
                need_emu += est_line_count(ptxt, sh.width, size) * size * ls * EMU_PER_PT
            if need_emu > sh.height * 1.35:
                print(f"  [可能溢出] P{idx:02d} {txt[:34]!r}")
                print(f"      估算需 {int(need_emu)}，框高 {sh.height}")
                bad += 1
    print("\n疑似溢出 %d 处（判据 1.35 倍，仅供参考）" % bad)
    return 0


if __name__ == "__main__":
    sys.exit(main())
